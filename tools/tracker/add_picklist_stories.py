"""Record the pick lists and postal code validation in the tracker.

SCR-137 Standard Pick Lists & Address Checks is not one screen: it is the
shared vocabulary and address rules eleven forms now draw on, recorded as
stories under a screen id the way SCR-136 records tenant isolation, so the work
is signed off in one place rather than scattered across the screens it touched.
"""

from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

import openpyxl

# Beside this script, not in the working directory: the tracker and the
# scripts that edit it live together in tools/tracker, and a relative
# name would quietly create or miss a copy wherever the script is run from.
FILE = str(Path(__file__).resolve().with_name("Chirala_Bay_Resort_PMS_Development_Tracker (2).xlsx"))
SHEET = "User Stories & Acceptance"

RANK, SEQ, SID, TITLE, EPIC, MODULE, SUBMODULE, SCREEN, SCREEN_NAME = 1, 2, 3, 4, 5, 6, 7, 8, 9
TYPE, PERSONA, STORY, VALUE, PRE, AC, NEG, RULES, DEPS, API = 10, 11, 12, 13, 14, 15, 16, 17, 18, 19
PRIORITY, RELEASE, POINTS, SPRINT, STATUS, PROGRESS, ASSIGNEE, QA, DOR, DOD, NOTES = (
    20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30)

DATE = "14 Sep 2026"

COMMON = dict(epic="EPIC-19", module="Administration", submodule="Data Quality",
              screen="SCR-137", screen_name="Standard Pick Lists & Address Checks",
              deps="SCR-005, SCR-009, SCR-026, SCR-059, Onboarding", priority="Should Have",
              release="MVP", sprint="S05-S07")

STORIES = [
    dict(sid="US-137-01", type_="Primary Workflow", points=5, status="Done",
         persona="Front Office Manager",
         title="Choose country, nationality, state and city instead of typing them",
         story=("As a front office manager, I want country, nationality and state as dropdowns and "
                "city suggestions for the chosen state, so that guest addresses are spelt one way"),
         value=("Reports and GST treatment group by these values; the guest table already held "
                "'Telangana', 'AP' and 'AndraPradash' for two states."),
         pre="Guest Check-In, New Reservation, Commercial Accounts, Invoice Settings or onboarding open.",
         ac=("Given an address form, when the country is India (the default) then State is a "
             "dropdown of the 36 states and union territories; for another country it is free "
             "text. Country and nationality are dropdowns of every country, searchable by typing. "
             "City accepts any value and suggests the chosen state's cities."),
         neg=("A stored value that is not in a list is shown as '(not in list)' and saved "
              "unchanged until someone chooses otherwise, so no existing address is erased by "
              "opening and saving a form."),
         api="frontend lib/options.ts (shared lists), components/ListSelect.tsx (ListSelect, StateField, CityInput).",
         notes=(f"BUILT {DATE}. One list file replaces the per-screen copies (onboarding had eight "
                "countries). Verified in the browser on Guest Check-In and Invoice Settings; "
                "Invoice Settings opened showing 'Andra Pradash (not in list)' rather than blank.")),
    dict(sid="US-137-02", type_="Primary Workflow", points=3, status="Done",
         persona="Finance Manager",
         title="Fill in the GST state code from the state",
         story=("As a finance manager, I want the GST state code filled in when I choose a state, so "
                "that CGST+SGST versus IGST is not decided by a hand-typed number"),
         value="The code decides the tax split on every invoice; typing it was the usual mistake.",
         pre="Invoice Settings, Commercial Accounts or onboarding billing.",
         ac=("Given an Indian state is chosen, then the GST state code is set to that state's code "
             "(e.g. Andhra Pradesh 37, Telangana 36) and can still be edited."),
         neg="A state with no GST code leaves the existing code alone.",
         api="gstCodeForState() over the GST_STATE_CODES map shared with GSTIN checking.",
         notes=(f"BUILT {DATE}. DATA TO REVIEW: Lekhana's Invoice Settings carry state 'Andra "
                "Pradash' with code 36 (Telangana) and a GSTIN beginning 36, for an address in "
                "Chirala, Andhra Pradesh (37). Left unchanged; the property must confirm where it "
                "is registered.")),
    dict(sid="US-137-03", type_="Primary Workflow", points=5, status="Done",
         persona="Property Admin",
         title="Set up rooms from one bed and view vocabulary and the Buildings & Floors master",
         story=("As a property admin, I want bed setup, view, building and floor chosen from the "
                "same lists everywhere, so that a room type made in onboarding opens correctly later"),
         value=("Onboarding stored 'King' while Room Types offered only '1 King Bed', so those types "
                "opened with a blank bed setup; rooms typed a building name unlinked to the estate."),
         pre="Buildings & Floors set up for the property.",
         ac=("Given the Room Editor, when buildings exist then Building and Floor are dropdowns from "
             "Buildings & Floors and saving sends the floor id, which sets both ids and names. Bed "
             "setup and view use one list on Room Editor, Room Types and onboarding."),
         neg=("A property with no buildings keeps the typed fields; a room whose stored building is "
              "not linked shows it as '(not linked)'; choosing a new building clears the floor."),
         api="GET /booking/buildings (with floors); PUT/POST /booking/rooms floor_id.",
         notes=(f"BUILT {DATE}. Verified on Edit Room 101: Building 'Lekhana Resort' and Floor '1st' "
                "selected from the master.")),
    dict(sid="US-137-04", type_="Primary Workflow", points=2, status="Done",
         persona="Property Admin",
         title="Pick timezone, currency and purpose of stay from lists",
         story=("As a property admin, I want timezone, currency and purpose of stay as dropdowns, "
                "so that they are valid values rather than free text"),
         value="A mistyped timezone shifts the business date; purpose of stay feeds segmentation.",
         pre="Property Settings, New Reservation or Reservation Detail.",
         ac=("Given Property Settings, then Timezone lists IANA zones and Currency lists codes with "
             "symbol and name; given a reservation, Purpose of Stay offers Leisure, Business, "
             "Wedding / Event, Conference, Family visit, Pilgrimage, Medical, Transit, Other."),
         neg="An unlisted stored timezone or currency stays selectable.",
         api="PUT property settings (unchanged); reservation purpose_of_stay (unchanged).",
         notes=f"BUILT {DATE}. Verified Property Settings shows Asia/Kolkata and 'INR (₹) · Indian Rupee'."),
    dict(sid="US-137-05", type_="Security / Exceptions", points=3, status="Done",
         persona="Front Office Manager",
         title="Refuse a malformed postal code on screen and at the API",
         story=("As a front office manager, I want a wrong PIN code caught as I type it, so that "
                "registration cards and invoices carry an address a letter could reach"),
         value="Guest records already held '34536' and '4242345' as Indian PIN codes.",
         pre="Any form with a postal code.",
         ac=("Given country India, when the postal code is not 6 digits starting 1-9 then the field "
             "shows 'A PIN code is 6 digits and does not start with 0.', only digits can be typed, "
             "and save / check-in / create reservation stay blocked. Other countries need 2-10 "
             "letters, digits, spaces or hyphens. The services refuse the same with 422 on guest "
             "save and check-in, guest create, commercial accounts, invoice settings and "
             "onboarding property."),
         neg=("An empty postal code is allowed everywhere. An API call with no country gets the "
              "shape check only, so a foreign code sent without a country is not refused."),
         api=("chirala_common.postal.problem; booking-core checkin_routes, routes (create_guest), "
              "account_routes; finance invoice_routes; iam onboarding_routes."),
         notes=(f"BUILT AND VERIFIED {DATE}. 22 unit tests (tests/test_postal_codes.py). Against the "
                "rebuilt services, PIN 34536 with country India was refused 422 by guest create, "
                "check-in guest save, commercial account create, invoice settings and onboarding "
                "property, each before any write; in the browser the field dropped a typed letter "
                "and space, showed the message and disabled Save Settings. Existing guests with "
                "malformed codes (34536, 4242345) are not altered; the desk is asked to correct "
                "them the next time the guest is saved.")),
]


def main() -> int:
    if any(Path(FILE).parent.glob("~$Chirala_Bay_Resort_PMS_Development_Tracker*")):
        print("The tracker is open in Excel; close it first.")
        return 1
    wb = openpyxl.load_workbook(FILE)
    ws = wb[SHEET]
    last = ws.max_row
    ids = {ws.cell(row=r, column=SID).value: r for r in range(5, last + 1)}
    if "US-137-01" in ids:
        print("SCR-137 stories already present; nothing written")
        return 0

    backup = FILE.replace(".xlsx", ".bak.20260914h.xlsx")
    n = 0
    while Path(backup).exists():
        n += 1
        backup = FILE.replace(".xlsx", f".bak.20260914h{n}.xlsx")
    shutil.copyfile(FILE, backup)
    print("backup:", backup)

    src = ids["US-136-01"]
    seq = max(int(v) for r in range(5, last + 1)
              if isinstance(v := ws.cell(row=r, column=SEQ).value, (int, float)))
    progress_template = next((ws.cell(row=r, column=PROGRESS).value for r in range(5, last + 1)
                              if str(ws.cell(row=r, column=PROGRESS).value or "").startswith("=IF(")), None)

    row = last + 1
    for st in STORIES:
        seq += 1
        for col in (RULES, DOR, DOD):
            ws.cell(row=row, column=col).value = ws.cell(row=src, column=col).value
        values = {
            RANK: 1, SEQ: seq, SID: st["sid"], TITLE: st["title"], EPIC: COMMON["epic"],
            MODULE: COMMON["module"], SUBMODULE: COMMON["submodule"], SCREEN: COMMON["screen"],
            SCREEN_NAME: COMMON["screen_name"], TYPE: st["type_"], PERSONA: st["persona"],
            STORY: st["story"], VALUE: st["value"], PRE: st["pre"], AC: st["ac"], NEG: st["neg"],
            DEPS: COMMON["deps"], API: st["api"], PRIORITY: COMMON["priority"],
            RELEASE: COMMON["release"], POINTS: st["points"], SPRINT: COMMON["sprint"],
            STATUS: st["status"], ASSIGNEE: "None", QA: f"TC-137-{chr(64 + int(st['sid'][-1]))}",
            NOTES: st["notes"],
        }
        for col, val in values.items():
            ws.cell(row=row, column=col).value = val
        ws.cell(row=row, column=PROGRESS).value = (
            re.sub(r"X\d+", f"X{row}", progress_template) if progress_template
            else {"Done": 1, "QA": 0.9}.get(st["status"], 0))
        print(f"  {st['sid']:<10} {st['status']:<5} row {row}  {st['title'][:55]}")
        row += 1

    wb.save(FILE)
    print(f"\n{len(STORIES)} stories written at rows {last + 1}-{row - 1}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
