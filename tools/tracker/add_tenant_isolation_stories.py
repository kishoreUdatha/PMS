"""Record database-level tenant isolation in the tracker, and close US-132-03.

SCR-136 Tenant Data Isolation is not a screen a user opens; it is the platform
guarantee every screen relies on, recorded the way the tracker records other
security work -- as stories under a screen id -- so it is signed off somewhere
rather than nowhere.

US-132-03 (Back Office Reports security) was held at QA because tenant refusal
had been proved at the endpoint but a mixed-tenant read had not been
demonstrated. It has now: every other tenant's data was altered inside a
rolled-back transaction and all 39 reports for Lekhana came out byte-identical.
"""

from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

import openpyxl

# Beside this script, not in the working directory: the tracker and the
# scripts that edit it live together in tools/tracker, and a relative
# name would quietly create or miss a copy wherever the script is run from.
FILE = str(Path(__file__).resolve().with_name("Chirala_Bay_Resort_PMS_Development_Tracker (2).xlsx"))
SHEET = "User Stories & Acceptance"

RANK, SEQ, SID, TITLE, EPIC, MODULE, SUBMODULE, SCREEN, SCREEN_NAME = 1, 2, 3, 4, 5, 6, 7, 8, 9
TYPE, PERSONA, STORY, VALUE, PRE, AC, NEG, RULES, DEPS, API = 10, 11, 12, 13, 14, 15, 16, 17, 18, 19
PRIORITY, RELEASE, POINTS, SPRINT, STATUS, PROGRESS, ASSIGNEE, QA, DOR, DOD, NOTES = (
    20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30)

DATE = "14 Sep 2026"

COMMON = dict(epic="EPIC-19", module="Administration", submodule="Tenant Isolation",
              screen="SCR-136", screen_name="Tenant Data Isolation",
              persona="Platform Operator", deps="All services", priority="Must Have",
              release="MVP", sprint="S05-S07")

STORIES = [
    dict(sid="US-136-01", type_="Security / Exceptions", points=13, status="Done",
         title="Enforce tenant isolation in the database itself",
         story=("As a platform operator, I want the database to refuse one tenant's data to "
                "another, so that a query that forgets its tenant filter cannot leak it"),
         value=("A multi-tenant SaaS promise that does not rest on every developer remembering "
                "a WHERE clause. Required by schema blueprint section 12."),
         pre="Services deployed through docker-compose with the runtime login configured.",
         ac=("Given any tenant-owned table, when a transaction bound to organisation A queries it "
             "with no filter at all, then only A's rows are returned; with no tenant bound, no "
             "rows; an insert or update naming organisation B is refused. Every application table "
             "(106) has row-level security enabled, forced, and a policy; the services connect as "
             "a login that is not superuser, cannot bypass RLS and owns no tables."),
         neg=("A migration refuses to finish if a table is left without a policy or a view exists "
              "that would read past row security; a service configured to connect as the schema "
              "owner fails a static test; the tenant setting is transaction-local and never "
              "carries to the next request on a pooled connection."),
         api=("Runtime login pms_app (chirala_common.runtime_role, run before each service starts); "
              "iam 0025_iam_rls, booking-core 0047_tenant_rls and 0048_outbox_rls, finance "
              "0028_finance_rls; tenancy.org_visible / system_context / identity_user_id."),
         notes=(f"BUILT AND VERIFIED {DATE} in four stages. Stage 1: runtime login and per-request "
                "tenant binding. Stage 2: finance (28 tables). Stage 3: booking, property, "
                "operations, guest and channel data (54 tables). Stage 4: identity, access and "
                "audit (23) and the event outbox (1). 30 database tests prove it from the runtime "
                "login, including queries with no WHERE clause, writes into another tenant, and "
                "that the identity function is SECURITY DEFINER with a pinned search path and no "
                "PUBLIC execute. The mixed-tenant leak test altered every other tenant's data and "
                "all 39 Lekhana reports came out identical. NOT COVERED: a service run straight "
                "from the host uses the owner URL in .env and bypasses RLS; production needs a "
                "strong APP_DB_PASSWORD.")),
    dict(sid="US-136-02", type_="Primary Workflow", points=8, status="Done",
         title="Give every unauthenticated and background path a narrow tenant context",
         story=("As a platform operator, I want the paths that run without a logged-in user to "
                "see only what their job needs, so that isolation holds outside ordinary requests"),
         value="The places isolation usually breaks -- jobs, webhooks, sign-in -- are the ones pinned down.",
         pre="US-136-01.",
         ac=("Given caller resolution, it reads only that subject's own identity rows and is "
             "narrowed to the caller's tenant before any handler runs; the public booking engine "
             "reads only the property in its URL, then binds that tenant; sign-in, sign-up and the "
             "platform console run in named system context; the hold reaper runs in system "
             "context; night audit, channel push and provisioning list work in system context and "
             "process each property as its own tenant; payment and channel webhooks use system "
             "context only until the tenant is identified."),
         neg=("Every system-context use states and logs its reason; binding a tenant switches "
              "system context off; the event outbox accepts inserts from any bound transaction "
              "but only system context can read it."),
         api=("chirala_common.db: bind_tenant_context, system_context, identity_context, "
              "property_code_context; permission dependencies bind the caller's organisation "
              "before the tenancy guard and the vetted property after it."),
         notes=(f"Verified {DATE}: a simulated Channex booking wrote its guest and hold into the "
                "right tenant and ended out of system context; channel push saw only its own "
                "tenant per link; the night-audit sweep saw every property while each audit saw "
                "one; sign-up seeded amenities invisible to other tenants; /me, sign-in refusal and "
                "forgot-password work; a GET smoke over 88 operations in three services returned "
                "no 500.")),
    dict(sid="US-136-03", type_="View / Search", points=2, status="Done",
         title="Platform console lists answer with no search filter",
         story=("As a platform operator, I want the organisations, properties and audit lists to "
                "open without typing a search term"),
         value="The console's own landing lists were unusable.",
         pre="Platform administrator.",
         ac=("Given no search term or status filter, when the organisations, properties or audit "
             "list is opened, then it returns every row the operator may see instead of an error."),
         neg="Optional filters are typed explicitly, so an empty filter is NULL rather than an untyped parameter.",
         api="GET /api/iam/platform/organizations, /properties, /audit.",
         notes=(f"DEFECT FOUND {DATE} while verifying tenant isolation, and not caused by it: the "
                "three lists failed identically with row security bypassed, because the SQL used "
                "':q IS NULL' with no type and Postgres refused an empty filter ('could not "
                "determine data type of parameter $1'). Fixed the same afternoon by a concurrent "
                "change to platform_routes.py (14:19, typed NULL checks), not by the isolation "
                "work; verified after the iam rebuild: organisations, properties and audit all "
                "answer 200 and the organisations list shows every tenant (4 of 4). /platform/users "
                "still requires a search term of at least three characters, by design.")),
]


def main() -> int:
    if any(Path(FILE).parent.glob("~$Chirala_Bay_Resort_PMS_Development_Tracker*")):
        print("The tracker is open in Excel; close it first.")
        return 1
    wb = openpyxl.load_workbook(FILE)
    ws = wb[SHEET]
    last = ws.max_row
    ids = {ws.cell(row=r, column=SID).value: r for r in range(5, last + 1)}
    if "US-136-01" in ids:
        print("SCR-136 stories already present; nothing written")
        return 0

    backup = FILE.replace(".xlsx", ".bak.20260914g.xlsx")
    n = 0
    while Path(backup).exists():
        n += 1
        backup = FILE.replace(".xlsx", f".bak.20260914g{n}.xlsx")
    shutil.copyfile(FILE, backup)
    print("backup:", backup)

    src = ids["US-131-03"]
    seq = max(int(v) for r in range(5, last + 1)
              if isinstance(v := ws.cell(row=r, column=SEQ).value, (int, float)))
    progress_template = next((ws.cell(row=r, column=PROGRESS).value for r in range(5, last + 1)
                              if str(ws.cell(row=r, column=PROGRESS).value or "").startswith("=IF(")), None)

    row = last + 1
    for st in STORIES:
        seq += 1
        for col in (RULES, DOR, DOD):
            ws.cell(row=row, column=col).value = ws.cell(row=src, column=col).value
        values = {
            RANK: 1, SEQ: seq, SID: st["sid"], TITLE: st["title"], EPIC: COMMON["epic"],
            MODULE: COMMON["module"], SUBMODULE: COMMON["submodule"], SCREEN: COMMON["screen"],
            SCREEN_NAME: COMMON["screen_name"], TYPE: st["type_"], PERSONA: COMMON["persona"],
            STORY: st["story"], VALUE: st["value"], PRE: st["pre"], AC: st["ac"], NEG: st["neg"],
            DEPS: COMMON["deps"], API: st["api"], PRIORITY: COMMON["priority"],
            RELEASE: COMMON["release"], POINTS: st["points"], SPRINT: COMMON["sprint"],
            STATUS: st["status"], ASSIGNEE: "None", QA: f"TC-136-{chr(64 + int(st['sid'][-1]))}",
            NOTES: st["notes"],
        }
        for col, val in values.items():
            ws.cell(row=row, column=col).value = val
        ws.cell(row=row, column=PROGRESS).value = (
            re.sub(r"X\d+", f"X{row}", progress_template) if progress_template
            else {"Done": 1, "QA": 0.9}.get(st["status"], 0))
        print(f"  {st['sid']:<10} {st['status']:<5} row {row}  {st['title'][:55]}")
        row += 1

    r = ids.get("US-132-03")
    if r:
        before = ws.cell(row=r, column=STATUS).value
        ws.cell(row=r, column=STATUS).value = "Done"
        if not str(ws.cell(row=r, column=PROGRESS).value or "").startswith("=IF("):
            ws.cell(row=r, column=PROGRESS).value = 1
        note = (f"[{DATE}] QA -> Done. The missing demonstration is done: inside a rolled-back "
                "transaction every other tenant's guests were renamed, their money multiplied by "
                "1000 and moved into Lekhana's dates, and foreign invoices, blocks, vouchers, "
                "owners and work orders added; all 39 Lekhana reports were byte-identical and the "
                "marker appeared in none, while the other tenants' own reports did change. The "
                "database now enforces the same boundary (SCR-136).")
        cell = ws.cell(row=r, column=NOTES)
        if "QA -> Done" not in str(cell.value or ""):
            cell.value = f"{cell.value}\n\n{note}" if cell.value else note
        print(f"  US-132-03  {before} -> Done")

    wb.save(FILE)
    print(f"\n{len(STORIES)} stories written at rows {last + 1}-{row - 1}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
