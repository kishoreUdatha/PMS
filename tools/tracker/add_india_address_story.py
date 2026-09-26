"""Add US-137-06: address rules that hold for every state in India.

The platform's properties can be anywhere in India, so the pick lists and
checks recorded under SCR-137 were widened from a regional first cut to all 36
states and union territories, and stored state names were cleaned for every
tenant.
"""

from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

import openpyxl

# Beside this script, not in the working directory: the tracker and the
# scripts that edit it live together in tools/tracker, and a relative
# name would quietly create or miss a copy wherever the script is run from.
FILE = str(Path(__file__).resolve().with_name("Chirala_Bay_Resort_PMS_Development_Tracker (2).xlsx"))
SHEET = "User Stories & Acceptance"

RANK, SEQ, SID, TITLE, EPIC, MODULE, SUBMODULE, SCREEN, SCREEN_NAME = 1, 2, 3, 4, 5, 6, 7, 8, 9
TYPE, PERSONA, STORY_COL, VALUE, PRE, AC, NEG, RULES, DEPS, API = 10, 11, 12, 13, 14, 15, 16, 17, 18, 19
PRIORITY, RELEASE, POINTS, SPRINT, STATUS, PROGRESS, ASSIGNEE, QA, DOR, DOD, NOTES = (
    20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30)

DATE = "14 Sep 2026"

STORY = dict(
    sid="US-137-06", type_="Security / Exceptions", points=5, status="Done",
    persona="Property Admin",
    title="Apply state, PIN code and GST code rules across all of India",
    story=("As a property admin anywhere in India, I want state names, PIN codes and GST state "
           "codes checked the same way for every state, so that no tenant's addresses depend on "
           "which region the product was first built for"),
    value=("A multi-tenant platform for all of India: 'AP', 'TN', 'Orissa' or 'AndraPradash' must "
           "all resolve to one jurisdiction, and a PIN or GST code from another state must be "
           "caught wherever the property is."),
    pre="US-137-01 to US-137-05.",
    ac=("Given any of the 36 states and union territories: City suggests that state's towns; a "
        "stored abbreviation, former name or close misspelling opens as the official state and "
        "the services store the official name on save (guest create, check-in, commercial "
        "accounts, invoice settings, onboarding property); a GST state code that belongs to a "
        "different state is shown under the field, blocks Save, and is refused with 422; a valid "
        "PIN code whose first two digits are issued in other states shows an amber warning."),
    neg=("The PIN/state check only warns, because postal zones cross state lines (Bihar/Jharkhand, "
         "UP/Uttarakhand, the smaller territories) and Army Postal Service codes (90-99) belong to "
         "no state. Regions outside India are never rewritten. GSTINs, GST codes and postal codes "
         "already stored are not changed by the clean-up; the screens flag them for the property."),
    api=("chirala_common.india (normalise_state, state_code_problem); frontend lib/options.ts "
         "(normaliseState, stateCodeProblem, pinStateWarning, CITIES_BY_STATE for all 36)."),
)


def main(notes: str) -> int:
    if any(Path(FILE).parent.glob("~$Chirala_Bay_Resort_PMS_Development_Tracker*")):
        print("The tracker is open in Excel; close it first.")
        return 1
    wb = openpyxl.load_workbook(FILE)
    ws = wb[SHEET]
    last = ws.max_row
    ids = {ws.cell(row=r, column=SID).value: r for r in range(5, last + 1)}
    if STORY["sid"] in ids:
        print(f"{STORY['sid']} already present; nothing written")
        return 0
    src = ids["US-137-05"]

    backup = FILE.replace(".xlsx", ".bak.20260914i.xlsx")
    n = 0
    while Path(backup).exists():
        n += 1
        backup = FILE.replace(".xlsx", f".bak.20260914i{n}.xlsx")
    shutil.copyfile(FILE, backup)
    print("backup:", backup)

    seq = max(int(v) for r in range(5, last + 1)
              if isinstance(v := ws.cell(row=r, column=SEQ).value, (int, float)))
    row = last + 1
    for col in (RULES, DOR, DOD, EPIC, MODULE, SUBMODULE, SCREEN, SCREEN_NAME, DEPS, PRIORITY,
                RELEASE, SPRINT):
        ws.cell(row=row, column=col).value = ws.cell(row=src, column=col).value
    values = {
        RANK: 1, SEQ: seq + 1, SID: STORY["sid"], TITLE: STORY["title"], TYPE: STORY["type_"],
        PERSONA: STORY["persona"], STORY_COL: STORY["story"], VALUE: STORY["value"],
        PRE: STORY["pre"], AC: STORY["ac"], NEG: STORY["neg"], API: STORY["api"],
        POINTS: STORY["points"], STATUS: STORY["status"], ASSIGNEE: "None", QA: "TC-137-F",
        NOTES: notes,
    }
    for col, val in values.items():
        ws.cell(row=row, column=col).value = val
    template = ws.cell(row=src, column=PROGRESS).value
    ws.cell(row=row, column=PROGRESS).value = (
        re.sub(r"X\d+", f"X{row}", template) if str(template or "").startswith("=IF(") else 1)

    # US-137-01 described a regional city list; say it now covers the country.
    r1 = ids.get("US-137-01")
    if r1:
        cell = ws.cell(row=r1, column=NOTES)
        add = (f"[{DATE}] Widened to all of India (US-137-06): city suggestions for all 36 states "
               "and union territories, none favoured.")
        if "US-137-06" not in str(cell.value or ""):
            cell.value = f"{cell.value}\n\n{add}" if cell.value else add

    wb.save(FILE)
    print(f"  {STORY['sid']}  Done  row {row}  {STORY['title']}")
    return 0


if __name__ == "__main__":
    notes_file = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    notes = notes_file.read_text(encoding="utf-8") if notes_file else f"BUILT {DATE}."
    sys.exit(main(notes))
