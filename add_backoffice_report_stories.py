"""Add the back-office reports and the three screens built for them to the tracker.

SCR-132 Back Office Reports  -- the /reports catalog and one report screen
                                 rendering all 39 reports of the Blue Way design
SCR-133 Work Orders           -- maintenance jobs (feeds report 38)
SCR-134 Expense Vouchers      -- raise / approve / pay (feeds report 16)
SCR-135 Unit Owners           -- owners and unit contracts (feeds report 27)

Statuses are written as verified, not as hoped. The reports and the three API
surfaces were exercised on 14 Sep 2026 against live Lekhana data and, for the
reports whose tables were empty, against seeded data inside a transaction
that was rolled back. Security stories stay QA: tenant refusal is proved at
the endpoint (403) but a mixed-tenant read has not been demonstrated, the same
standard US-131-03 was held to.

US-027-01 (Reports Dashboard) gets a note rather than a status change: /reports
is now a catalog of reports, which is close to but not the mockup-027 screen.
"""

from __future__ import annotations

import re
import shutil
import sys
from collections import Counter
from pathlib import Path

import openpyxl

FILE = "Chirala_Bay_Resort_PMS_Development_Tracker (2).xlsx"
SHEET = "User Stories & Acceptance"

# Column numbers (1-based) in the stories sheet.
RANK, SEQ, SID, TITLE, EPIC, MODULE, SUBMODULE, SCREEN, SCREEN_NAME = 1, 2, 3, 4, 5, 6, 7, 8, 9
TYPE, PERSONA, STORY, VALUE, PRE, AC, NEG, RULES, DEPS, API = 10, 11, 12, 13, 14, 15, 16, 17, 18, 19
PRIORITY, RELEASE, POINTS, SPRINT, STATUS, PROGRESS, ASSIGNEE, QA, DOR, DOD, NOTES = (
    20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30)

VERIFIED = "14 Sep 2026"

SCREENS = [
    dict(
        screen="SCR-132", name="Back Office Reports", module="Analytics",
        submodule="Management Reports", epic_from_module="Analytics",
        persona="Finance Manager", deps="Finance, Reservations, Housekeeping",
        stories=[
            dict(type_="View / Search", points=8, status="Done",
                 title="View and find information in Back Office Reports",
                 story=("As a Finance Manager, I want one catalog of every back-office "
                        "report, searchable and grouped by category, and to open any "
                        "report on this property's live data"),
                 value="Replaces 39 separate exports with one place that reads the same records the operational screens do.",
                 pre="Authenticated user; active property; reports.view permission.",
                 ac=("Given I open /reports, then all 39 reports of the back-office design and the "
                     "Hotel Ledger are listed by category with search and favourites; when I open "
                     "one, then it runs for the property's open business date by default and shows "
                     "metric tiles, a table with totals and the basis of every figure."),
                 neg=("A report with no data says why (e.g. no invoice issued yet) rather than showing "
                      "zeroes; windows over 366 days and an end date before the start are refused."),
                 api=("GET /api/finance/reports/backoffice (catalog); GET "
                      "/api/finance/reports/backoffice/{slug} (property_id, date_from, date_to, "
                      "report parameters). One engine: each report is a definition plus one query."),
                 notes=(f"BUILT AND VERIFIED {VERIFIED}. All 39 reports live. Cross-checked on Lekhana "
                        "1-14 Sep: Daily Revenue = Detail Revenue (61,000); Receipt Summary = Receipt "
                        "Detail (1,20,224); Transaction Detail debits/credits = Hotel Ledger; Folio List "
                        "balance = Guest Ledger closing; room revenue identical across the four room "
                        "reports (22,000). Reports with empty tables verified against seeded data in a "
                        "rolled-back transaction: ageing buckets, GST CGST/SGST split and credit note, "
                        "10% commission, out-of-order nights removed from occupancy, half-board covers, "
                        "owner statement arithmetic, work-order overdue as of a date. "
                        "DECISIONS: Police Inquiry List masks ID numbers to the last four characters; "
                        "Cashier Sales omits sales before tax (charges carry no cashier); card fees are "
                        "not recorded so card amounts are gross; half board counted as breakfast + "
                        "dinner; complimentary stays are derived (zero rate, charges adjusted off, "
                        "complimentary upgrade) because no complimentary flag exists.")),
            dict(type_="Primary Workflow", points=5, status="Done",
                 title="Complete the primary Back Office Reports workflow",
                 story=("As a Finance Manager, I want to change the dates, narrow by the report's own "
                        "filters, choose columns, inspect a record and export or print what I see"),
                 value="Month-end and night-audit checks become a filtered, exportable statement instead of spreadsheet work.",
                 pre="View access confirmed.",
                 ac=("Given a report is open, when I set dates and filters and run it, then tiles, "
                     "table and totals describe exactly the rows shown; the CSV contains the visible "
                     "columns with figures as numbers; a row opens a drawer with every field and "
                     "links to its folio and reservation."),
                 neg="Column filters never change the server query; totals are never computed over hidden rows.",
                 api="Same endpoints; CSV and print are client-side from the payload.",
                 notes=f"Verified in the browser {VERIFIED}; no console errors."),
            dict(type_="Security / Exceptions", points=3, status="QA",
                 title="Control permissions, overrides and exceptions in Back Office Reports",
                 story=("As a department supervisor, I want reports refused to callers without "
                        "reporting rights and for other tenants' properties"),
                 value="Reports expose the whole property's money and guests' identity documents.",
                 pre="Roles, permissions and tenant scoping configured.",
                 ac="Given a user lacks reports.view, or names another organisation's property, then the request is refused.",
                 neg="Unbuilt report slugs return 404; Hotel Ledger is not reachable through the generic route.",
                 api="require_permission('reports','view') + assert_property_in_org on both endpoints.",
                 notes=("QA, not Done: refusal is enforced by the same guards as US-131-03 and every "
                        "report is run per property in the integration test, but a mixed-tenant "
                        "window has not been demonstrated.")),
        ],
    ),
    dict(
        screen="SCR-133", name="Work Orders", module="Housekeeping",
        submodule="Maintenance", epic_from_module="Housekeeping",
        persona="Maintenance Supervisor", deps="Rooms, Users",
        stories=[
            dict(type_="View / Search", points=3, status="Done",
                 title="View and find information in Work Orders",
                 story="As a Maintenance Supervisor, I want to see open work by priority and due date, with overdue jobs counted",
                 value="Faults stop living on paper and in chat, and overdue work is visible every morning.",
                 pre="housekeeping.view permission.",
                 ac="Given I open Work Orders, then open jobs list most urgent first with counts per status and overdue.",
                 neg="Another tenant's property is refused 403.",
                 api="GET /api/booking/work-orders (property_id, status|active, priority, q).",
                 notes=f"Verified {VERIFIED} via API in a rolled-back transaction: ordering, counts, overdue, tenant refusal."),
            dict(type_="Primary Workflow", points=5, status="Done",
                 title="Complete the primary Work Orders workflow",
                 story="As a Maintenance Supervisor, I want to raise a job against a room or place, assign it and move it to completion",
                 value="Days-open on the Work Order List report is measured, not remembered.",
                 pre="housekeeping.create / housekeeping.edit.",
                 ac=("Given a job is raised, then it is numbered WO-1001 onwards per property; In progress "
                     "stamps started_at, Completed stamps completed_at and reopening clears it."),
                 neg="A job with neither room nor location, an unknown category or an assignee outside the organisation is refused.",
                 api="POST /api/booking/work-orders; PATCH /api/booking/work-orders/{id}. Table operations.work_orders (booking-core 0046).",
                 notes=f"Verified {VERIFIED}: numbering, stamps, validation. Feeds report 38."),
        ],
    ),
    dict(
        screen="SCR-134", name="Expense Vouchers", module="Finance",
        submodule="Payables", epic_from_module="Finance",
        persona="Finance Manager", deps="Rooms, Approvals",
        stories=[
            dict(type_="Primary Workflow", points=5, status="Done",
                 title="Complete the primary Expense Vouchers workflow",
                 story="As a Finance Manager, I want vouchers raised, approved or rejected, then paid, with who decided recorded",
                 value="Every rupee out that is not a refund has an approver and a record.",
                 pre="payments.create to raise; payments.approve to decide; payments.edit to pay or cancel.",
                 ac=("Given a voucher is raised it is pending approval and numbered EV-1001 onwards; only a "
                     "pending voucher can be edited; approval records the approver; only an approved "
                     "voucher can be paid."),
                 neg="Rejection without a reason, paying before approval and editing an approved voucher are refused.",
                 api=("GET/POST /api/finance/expenses; PATCH /api/finance/expenses/{id}; POST "
                      "/api/finance/expenses/{id}/decision. Table finance.expense_vouchers (finance 0027)."),
                 notes=(f"Verified {VERIFIED} via API in a rolled-back transaction. A defect found in that "
                        "run was fixed: the decision UPDATE used one bind as both a varchar column and a "
                        "text cast, which Postgres refuses. Feeds reports 16 and 27.")),
        ],
    ),
    dict(
        screen="SCR-135", name="Unit Owners", module="Finance",
        submodule="Owner Accounting", epic_from_module="Finance",
        persona="Property Owner", deps="Rooms, Expense Vouchers",
        stories=[
            dict(type_="Primary Workflow", points=5, status="Done",
                 title="Complete the primary Unit Owners workflow",
                 story="As a Resort Manager, I want to record unit owners and which rooms they hold on what management fee",
                 value="Owner statements can be produced from the ledger instead of a spreadsheet.",
                 pre="payments.view / payments.edit.",
                 ac="Given an owner and a contract are saved, then the room shows its current owner and the Owner Statement settles it.",
                 neg="Overlapping contracts on the same room and contracts ending before they start are refused.",
                 api=("GET/POST /api/finance/unit-owners; PATCH /api/finance/unit-owners/{id}; POST "
                      "/api/finance/unit-owners/{id}/contracts; PATCH /api/finance/unit-owner-contracts/{id}. "
                      "Tables finance.unit_owners, finance.unit_owner_contracts (finance 0027)."),
                 notes=(f"Verified {VERIFIED}: overlap refusal, date validation, current owner, and the "
                        "statement (16,000 revenue, 20% fee 3,200, 2,860 charges, 9,940 net). Payouts "
                        "to owners are not recorded yet.")),
        ],
    ),
]


def main() -> int:
    if any(Path(".").glob("~$Chirala_Bay_Resort_PMS_Development_Tracker*")):
        print("The tracker is open in Excel; close it first.")
        return 1
    wb = openpyxl.load_workbook(FILE)
    ws = wb[SHEET]
    last = ws.max_row
    existing = {ws.cell(row=r, column=SID).value for r in range(5, last + 1)}
    if "US-132-01" in existing:
        print("SCR-132 stories already present; nothing written")
        return 0

    backup = FILE.replace(".xlsx", ".bak.20260914f.xlsx")
    n = 0
    while Path(backup).exists():
        n += 1
        backup = FILE.replace(".xlsx", f".bak.20260914f{n}.xlsx")
    shutil.copyfile(FILE, backup)
    print("backup:", backup)

    src = next(r for r in range(5, last + 1) if ws.cell(row=r, column=SID).value == "US-131-01")
    seq = max(int(v) for r in range(5, last + 1)
              if isinstance(v := ws.cell(row=r, column=SEQ).value, (int, float)))
    template_progress = next(
        (ws.cell(row=r, column=PROGRESS).value for r in range(5, last + 1)
         if str(ws.cell(row=r, column=PROGRESS).value or "").startswith("=IF(")), None)

    def epic_for(module: str) -> str:
        counts = Counter(ws.cell(row=r, column=EPIC).value for r in range(5, last + 1)
                         if ws.cell(row=r, column=MODULE).value == module
                         and str(ws.cell(row=r, column=EPIC).value or "").startswith("EPIC-"))
        return counts.most_common(1)[0][0] if counts else "EPIC-16"

    row = last + 1
    for screen in SCREENS:
        epic = epic_for(screen["epic_from_module"])
        for i, st in enumerate(screen["stories"], start=1):
            seq += 1
            sid = f"US-{screen['screen'][4:]}-{i:02d}"
            for col in (RULES, DOR, DOD):
                ws.cell(row=row, column=col).value = ws.cell(row=src, column=col).value
            values = {
                RANK: 1, SEQ: seq, SID: sid, TITLE: st["title"], EPIC: epic,
                MODULE: screen["module"], SUBMODULE: screen["submodule"],
                SCREEN: screen["screen"], SCREEN_NAME: screen["name"],
                TYPE: st["type_"], PERSONA: screen["persona"], STORY: st["story"],
                VALUE: st["value"], PRE: st["pre"], AC: st["ac"], NEG: st["neg"],
                DEPS: screen["deps"], API: st["api"], PRIORITY: "Must Have",
                RELEASE: "MVP", POINTS: st["points"], SPRINT: "S05-S07",
                STATUS: st["status"], ASSIGNEE: "None", QA: f"TC-{screen['screen'][4:]}-{chr(64 + i)}",
                NOTES: st["notes"],
            }
            for col, val in values.items():
                ws.cell(row=row, column=col).value = val
            ws.cell(row=row, column=PROGRESS).value = (
                re.sub(r"X\d+", f"X{row}", template_progress) if template_progress
                else {"Done": 1, "QA": 0.9}.get(st["status"], 0))
            print(f"  {sid:<10} {st['status']:<5} row {row}  {st['title'][:55]}")
            row += 1

    for r in range(5, last + 1):
        if ws.cell(row=r, column=SID).value == "US-027-01":
            cell = ws.cell(row=r, column=NOTES)
            note = (f"[{VERIFIED}] /reports now renders the Back Office Reports catalog (SCR-132, "
                    "US-132-01..03): every report of the back-office design, searchable by category, "
                    "39 of them live. That is a catalog of reports, not the mockup-027 dashboard, so "
                    "this story's status is left unchanged.")
            if note[:14] not in str(cell.value or ""):
                cell.value = f"{cell.value}\n\n{note}" if cell.value else note
                print("  US-027-01  note added")
            break

    wb.save(FILE)
    print(f"\n{row - last - 1} stories written at rows {last + 1}-{row - 1}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
